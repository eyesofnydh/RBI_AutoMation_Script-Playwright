"""Builds RBI-Test-Cases.xlsx — run: python tools/build_test_cases.py <output.xlsx>"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else 'RBI-Test-Cases.xlsx'
BASE = 'https://stg-rbi.webc.in'

# ----------------------------------------------------------------------------------------------
# Test cases: (module, feature, page, scenario, preconditions, steps, data, expected, type, priority, automated_in)
# automated_in = spec file (Automated) or '' (Manual)
# ----------------------------------------------------------------------------------------------
H = 'tests/functional/header.spec.ts'
S = 'tests/functional/search.spec.ts'
AL = 'tests/functional/accessibility-language.spec.ts'
HP = 'tests/functional/homepage.spec.ts'
LS = 'tests/functional/listings.spec.ts'
TP = 'tests/functional/templates.spec.ts'
RS = 'tests/functional/responsive.spec.ts'
PG = 'tests/crawl/pages.spec.ts'
LK = 'tests/crawl/links.spec.ts'
AX = 'tests/crawl/a11y.spec.ts'
IN = 'tests/crawl/interactions.spec.ts'
XL = 'tests/functional/xpath-locators.spec.ts'

PRE_HOME = 'Staging login done; homepage open'
PRE_ANY = 'Staging login done'

cases = []
def add(module, feature, page, scenario, pre, steps, data, expected, typ, pri, auto=''):
    cases.append((module, feature, page, scenario, pre, steps, data, expected, typ, pri, auto))

# ============================== 1. ACCESS & ENVIRONMENT ==============================
M = 'Access & Environment'
add(M, 'Basic auth', '/', 'Site asks for credentials on first visit', 'Fresh incognito window', '1. Open the staging URL', '', 'Browser sign-in pop-up appears; page not shown until login', 'Security', 'P1')
add(M, 'Basic auth', '/', 'Valid credentials open the site', 'Fresh incognito window', '1. Open staging URL\n2. Enter valid username/password', 'Valid staging credentials', 'Homepage loads (HTTP 200)', 'Functional', 'P1', PG)
add(M, 'Basic auth', '/', 'NEGATIVE: wrong password is rejected', 'Fresh incognito window', '1. Open staging URL\n2. Enter valid user + wrong password', 'Wrong password', 'HTTP 401; pop-up shown again; no content leaked', 'Negative', 'P1')
add(M, 'Basic auth', '/', 'NEGATIVE: cancel the sign-in pop-up', 'Fresh incognito window', '1. Open staging URL\n2. Click Cancel', '', '401 page; no site content visible', 'Negative', 'P2')
add(M, 'HTTPS', '/', 'HTTP redirects to HTTPS', '', '1. Open http://stg-rbi.webc.in', '', 'Redirects to https://; valid certificate, padlock shown', 'Security', 'P1')
add(M, 'Security headers', '/', 'Security headers present', '', '1. Open DevTools → Network → document response headers', '', 'Strict-Transport-Security, X-Content-Type-Options: nosniff, X-Frame-Options/CSP frame-ancestors, Referrer-Policy present', 'Security', 'P2')
add(M, 'Clickjacking', '/', 'Site cannot be framed by another origin', '', '1. Create an HTML page with <iframe src="staging URL">\n2. Open it', '', 'Frame refuses to load (X-Frame-Options / CSP)', 'Security', 'P2')
add(M, 'Indexing', '/', 'Staging is not indexable by search engines', '', '1. Check robots meta / X-Robots-Tag on staging', '', 'noindex present on staging (production must NOT have it)', 'SEO', 'P2')

# ============================== 2. HEADER ==============================
M = 'Header'
add(M, 'Logo', 'All pages', 'Logo visible and links to homepage', PRE_ANY, '1. Open /press-releases\n2. Click RBI logo', '', 'Homepage opens', 'Functional', 'P1', H)
add(M, 'Logo', 'All pages', 'Logo image has alt text / accessible name', PRE_ANY, '1. Inspect logo link', '', 'Link has accessible name e.g. "Reserve Bank of India – Home"', 'Accessibility', 'P2')
add(M, 'Primary nav', 'All pages', 'Press Releases link opens listing and is marked current', PRE_ANY, '1. Click "Press Releases" in header', '', '/press-releases opens; link has aria-current="page"', 'Functional', 'P1', H)
add(M, 'Mega-menu', 'All pages', '"About RBI" mega-menu opens with links', PRE_HOME, '1. Click "About RBI"', '', 'Panel opens; aria-expanded=true; links visible', 'Functional', 'P1', H)
add(M, 'Mega-menu', 'All pages', '"Functions" mega-menu opens with links', PRE_HOME, '1. Click "Functions"', '', 'Panel opens; aria-expanded=true; links visible', 'Functional', 'P1', H)
add(M, 'Mega-menu', 'All pages', 'Every mega-menu link resolves (no 404)', PRE_HOME, '1. Open each mega-menu\n2. Visit every link', '', 'All links return 200. KNOWN DEFECT: /regulation → 404', 'Functional', 'P1', H)
add(M, 'Mega-menu', 'All pages', 'Only one mega-menu open at a time', PRE_HOME, '1. Open About RBI\n2. Click Functions', '', 'About RBI closes; Functions opens', 'UI', 'P2', H)
add(M, 'Mega-menu', 'All pages', 'Escape closes mega-menu and returns focus', PRE_HOME, '1. Open About RBI\n2. Press Esc', '', 'Menu closes; focus back on "About RBI"', 'Accessibility', 'P2', H)
add(M, 'Mega-menu', 'All pages', 'Click outside closes mega-menu', PRE_HOME, '1. Open Functions\n2. Click on page body', '', 'Menu closes', 'UI', 'P3')
add(M, 'Mega-menu', 'All pages', 'Keyboard: Tab through mega-menu links', PRE_HOME, '1. Tab to About RBI, press Enter\n2. Tab through links', '', 'Focus moves through links in order; visible focus ring', 'Accessibility', 'P2')
add(M, 'Mega-menu', 'All pages', 'EDGE: rapid open/close clicking', PRE_HOME, '1. Click About RBI 10 times quickly', '', 'Ends in consistent state; no flicker, no JS error', 'Edge', 'P3', IN)
add(M, 'Hamburger menu', 'All pages', 'Menu opens and lists all sections', PRE_HOME, '1. Click "Menu"', '', 'Full-screen/side menu shows sections; aria-expanded=true', 'Functional', 'P1', H)
add(M, 'Hamburger menu', 'All pages', 'Menu closes with close button and Esc', PRE_HOME, '1. Open Menu\n2. Click Close / press Esc', '', 'Menu closes; page scroll restored', 'Functional', 'P1', H)
add(M, 'Hamburger menu', 'All pages', 'All hamburger menu links resolve', PRE_HOME, '1. Open Menu\n2. Visit every link', '', 'All links return 200', 'Functional', 'P1', H)
add(M, 'Hamburger menu', 'All pages', 'Focus is trapped inside open menu', PRE_HOME, '1. Open Menu\n2. Press Tab repeatedly', '', 'Focus cycles inside menu; does not reach page behind', 'Accessibility', 'P2')
add(M, 'Hamburger menu', 'All pages', 'Page behind does not scroll while menu open', PRE_HOME, '1. Open Menu\n2. Scroll mouse wheel', '', 'Only the menu scrolls', 'UI', 'P3')
add(M, 'Skip link', 'All pages', 'First Tab stop is "Skip to main content"', PRE_HOME, '1. Load page\n2. Press Tab once', '', 'Skip link becomes visible and focused', 'Accessibility', 'P2', H)
add(M, 'Skip link', 'All pages', 'Skip link moves focus into main content', PRE_HOME, '1. Tab to skip link\n2. Press Enter', '', 'Focus lands on <main>; next Tab goes to first content link', 'Accessibility', 'P2', H)
add(M, 'Sticky header', 'All pages', 'Header stays reachable after scrolling', PRE_HOME, '1. Scroll to bottom', '', 'Header/menu button still visible (sticky) or reachable', 'UI', 'P3', H)
add(M, 'Header layout', 'All pages', 'Header consistent across all page templates', PRE_ANY, '1. Visit home, listing, detail, static, search pages', '', 'Same header layout, no shift, logo size consistent', 'UI', 'P2')

# ============================== 3. SEARCH ==============================
M = 'Search'
add(M, 'Header search', 'All pages', 'Search with valid keyword shows results', PRE_HOME, '1. Type keyword in header search\n2. Press Enter', 'repo rate', '/search?q=repo+rate opens; heading "Search result for “repo rate”"; relevant results', 'Functional', 'P1', S)
add(M, 'Header search', 'All pages', 'Search via Search button (not Enter)', PRE_HOME, '1. Type keyword\n2. Click magnifier button', 'inflation', 'Same results as pressing Enter', 'Functional', 'P2')
add(M, 'Results', '/search', 'Result count shown and numeric', PRE_ANY, '1. Search "repo rate"', 'repo rate', '"Show N results" shows a number > 0', 'Functional', 'P2', S)
add(M, 'Results', '/search', 'Clicking a result opens the right page', PRE_ANY, '1. Search\n2. Click first result', 'repo rate', 'Correct detail page opens with matching title', 'Functional', 'P1', S)
add(M, 'Results', '/search', 'Results paginate (page 2, next, previous)', PRE_ANY, '1. Search\n2. Click page 2, Next, Previous', 'repo rate', 'Results change per page; no duplicates across pages', 'Functional', 'P2', S)
add(M, 'Results', '/search', 'Items-per-page selector (5/10/20/50)', PRE_ANY, '1. Search\n2. Change items per page', 'repo rate', 'Number of results per page matches selection', 'Functional', 'P3')
add(M, 'Results', '/search', 'Function filter narrows results', PRE_ANY, '1. Search\n2. Tick a function (e.g. About Us)\n3. Apply', 'repo rate', 'Only results for that function; count updates', 'Functional', 'P2')
add(M, 'Results', '/search', 'Date filter (month / custom range) narrows results', PRE_ANY, '1. Search\n2. Pick month or custom date range', 'repo rate; Jan 2020–Dec 2020', 'Results only within range', 'Functional', 'P2')
add(M, 'Results', '/search', '"Clear all" removes query and filters', PRE_ANY, '1. Search with filters\n2. Click Clear all', '', 'Filters/query cleared', 'Functional', 'P2', S)
add(M, 'Results', '/search', 'Search term is highlighted in results', PRE_ANY, '1. Search', 'repo', 'Matches highlighted (if designed)', 'UI', 'P3')
add(M, 'Edge', '/search', 'EDGE: no-results message', PRE_ANY, '1. Search gibberish', 'zqxwvkjhg123', '"No results for …" message with suggestion; no error', 'Edge', 'P1', S)
add(M, 'Edge', '/search', 'EDGE: empty search', PRE_HOME, '1. Leave box empty\n2. Press Enter', '(empty)', 'Validation message or no action; no crash', 'Edge', 'P2', S)
add(M, 'Edge', '/search', 'EDGE: only spaces', PRE_HOME, '1. Enter spaces only\n2. Press Enter', '"     "', 'Treated as empty', 'Edge', 'P3', IN)
add(M, 'Edge', '/search', 'EDGE: max length 250 enforced', PRE_HOME, '1. Paste 400 characters', "'a' × 400", 'Input stops at 250 characters', 'Edge', 'P2', S)
add(M, 'Edge', '/search', 'EDGE: special characters', PRE_HOME, '1. Search special chars', '~!@#$%^&*()"<>', 'No error; handled safely', 'Edge', 'P2', S)
add(M, 'Edge', '/search', 'EDGE: Hindi / Unicode / emoji query', PRE_HOME, '1. Search Unicode', 'भारतीय रिज़र्व बैंक ₹ 😀', 'No error; relevant or no-results message', 'Edge', 'P2', IN)
add(M, 'Edge', '/search', 'EDGE: case-insensitive', PRE_HOME, '1. Search "REPO RATE" then "repo rate"', '', 'Same result count', 'Edge', 'P3')
add(M, 'Edge', '/search', 'EDGE: leading/trailing spaces trimmed', PRE_HOME, '1. Search "  repo rate  "', '', 'Same results as "repo rate"', 'Edge', 'P3')
add(M, 'Edge', '/search', 'EDGE: direct URL with query', PRE_ANY, '1. Open /search?q=inflation directly', '', 'Results load without using the box', 'Edge', 'P2')
add(M, 'Edge', '/search', 'EDGE: browser Back after search', PRE_HOME, '1. Search\n2. Open a result\n3. Press Back', '', 'Returns to same results page, same filters & page number', 'Edge', 'P2')
add(M, 'Security', '/search', 'SECURITY: XSS in search box', PRE_HOME, '1. Search script payload', '<script>alert(1)</script>', 'Text shown escaped; no alert', 'Security', 'P1', S)
add(M, 'Security', '/search', 'SECURITY: XSS in URL parameter', PRE_ANY, '1. Open /search?q=<img src=x onerror=alert(1)>', '', 'No alert; payload shown as text', 'Security', 'P1')
add(M, 'Security', '/search', 'SECURITY: SQL-like input', PRE_HOME, "1. Search SQL payload", "' OR '1'='1'; --", 'No 500 error; no DB error text', 'Security', 'P1', IN)
add(M, 'Mobile', '/search', 'Search from mobile menu', 'Mobile device / 375px', '1. Open Menu\n2. Use search box in menu', 'repo rate', 'Results load', 'Functional', 'P2', S)

# ============================== 4. ACCESSIBILITY TOOLBAR ==============================
M = 'Accessibility Toolbar'
add(M, 'Panel', 'All pages', 'Toolbar opens as dialog with 10 tools', PRE_HOME, '1. Click Accessibility icon', '', 'Dialog shows Light, Dark, Increase Text, Decrease Text, Reset Text, Text Spacing, Line Height, Hide Images, Big Cursor, Screen Reader, Reset All', 'Functional', 'P1', AL)
add(M, 'Panel', 'All pages', 'Close button and Esc close panel; focus returns', PRE_HOME, '1. Open panel\n2. Close / Esc', '', 'Panel closes; focus on Accessibility icon', 'Accessibility', 'P2', AL)
for tool, exp in [('Dark', 'Dark theme applied; text readable (contrast ≥ 4.5:1)'), ('Increase Text', 'Text grows (e.g. 16→16.8px); layout does not break'), ('Decrease Text', 'Text shrinks; stays readable'), ('Text Spacing', 'Letter/word spacing increases'), ('Line Height', 'Line height increases'), ('Hide Images', 'Images hidden; alt text / layout OK'), ('Big Cursor', 'Large cursor shown'), ('Screen Reader', 'Screen reader helper enabled (as designed)')]:
    add(M, tool, 'All pages', f'"{tool}" changes the page and Reset All restores it', PRE_HOME, f'1. Open panel\n2. Click "{tool}"\n3. Click Reset All', '', f'{exp}; Reset All restores default', 'Functional', 'P2', AL)
add(M, 'Text size', 'All pages', 'EDGE: Increase Text clicked 10 times', PRE_HOME, '1. Click Increase Text 10×', '', 'Stops at a max size; no overlapping text or horizontal scroll', 'Edge', 'P2')
add(M, 'Text size', 'All pages', 'EDGE: Decrease Text clicked 10 times', PRE_HOME, '1. Click Decrease Text 10×', '', 'Stops at a min size; still readable', 'Edge', 'P2')
add(M, 'Text size', 'All pages', 'Reset Text returns to default size', PRE_HOME, '1. Increase twice\n2. Reset Text', '', 'Font size back to default', 'Functional', 'P2', AL)
add(M, 'Theme', 'All pages', 'Light / Dark are mutually exclusive', PRE_HOME, '1. Click Dark\n2. Click Light', '', 'Only one pressed at a time', 'Functional', 'P2', AL)
add(M, 'Persistence', 'All pages', 'Settings persist across pages', PRE_HOME, '1. Enable Dark\n2. Navigate to another page', '', 'Dark still applied', 'Functional', 'P2', AL)
add(M, 'Persistence', 'All pages', 'Settings persist after browser refresh', PRE_HOME, '1. Enable Increase Text\n2. Refresh', '', 'Setting kept (if designed) or reset consistently', 'Edge', 'P3')
add(M, 'Combination', 'All pages', 'EDGE: all tools on at once', PRE_HOME, '1. Enable Dark + Increase Text ×3 + Text Spacing + Line Height + Hide Images', '', 'Page still usable; no overlaps; menus work', 'Edge', 'P2')
add(M, 'Dark theme', 'All pages', 'Dark mode on every template', PRE_ANY, '1. Enable Dark\n2. Visit home, listing, detail, static, search, offices', '', 'No white flashes, unreadable text or invisible icons', 'UI', 'P2')

# ============================== 5. LANGUAGE ==============================
M = 'Language'
add(M, 'Switch', 'All pages', 'Change language opens language choice', PRE_HOME, '1. Click "Change language" icon', '', 'Menu with हिन्दी / English. KNOWN DEFECT: nothing happens', 'Functional', 'P1', AL)
add(M, 'Switch', 'All pages', 'Switch to Hindi translates page', PRE_HOME, '1. Choose हिन्दी', '', 'Content in Hindi; <html lang="hi">; URL /hi/... KNOWN DEFECT: /hi → 404', 'Functional', 'P1', AL)
add(M, 'Switch', 'Inner pages', 'Hindi on an inner page keeps same page', PRE_ANY, '1. Open /press-releases\n2. Switch to Hindi', '', 'Hindi version of Press Releases (not homepage)', 'Functional', 'P2', AL)
add(M, 'Switch', 'All pages', 'Switch back to English', 'Hindi page open', '1. Choose English', '', 'English content; lang="en"', 'Functional', 'P1', AL)
add(M, 'Hindi UI', 'All pages', 'Devanagari renders correctly (no boxes)', 'Hindi page open', '1. Check headings, menus, footer', '', 'Correct font; no tofu boxes; no clipped matras', 'UI', 'P2')
add(M, 'Edge', 'All pages', 'EDGE: Hindi page with no translation', 'Hindi page open', '1. Open a page with no Hindi content', '', 'Friendly fallback (English or message), not 404', 'Edge', 'P2')

# ============================== 6. HOMEPAGE ==============================
M = 'Homepage'
P = '/'
add(M, 'Hero', P, 'Hero vision statement renders', PRE_HOME, '1. Load homepage', '', '"Continue excellence while enabling…" + "RBI\'s Vision Statement (Utkarsh 2029)"; lion image loads', 'UI', 'P1', HP)
add(M, 'Hero', P, 'Page has one (screen-reader) H1', PRE_HOME, '1. Inspect headings', '', 'Exactly one H1 "Reserve Bank of India"', 'Accessibility', 'P3', HP)
for hub, label, inner in [('notifications', 'Notifications', ''), ('citizen-s-centre', "Citizen's Centre", 'Services / Information / Useful Links'), ('research-publication', 'Publications & Research', 'Publications / Research / Statistics'), ('speeches-and-media', 'Speeches & Media', 'Speeches / …')]:
    add(M, 'Quick-access hub', P, f'"{label}" tile expands its panel', PRE_HOME, f'1. Click "{label}" tile', '', f'Panel opens (aria-expanded=true) with links{("; inner tabs " + inner) if inner else ""}', 'Functional', 'P1', HP)
    if inner:
        add(M, 'Quick-access hub', P, f'"{label}" inner tabs switch content', PRE_HOME, f'1. Open "{label}"\n2. Click each inner tab', '', 'Each tab shows its own panel; aria-selected updates', 'Functional', 'P2', HP)
add(M, 'Quick-access hub', P, 'Only one hub panel open at a time', PRE_HOME, '1. Open Notifications\n2. Open Citizen\'s Centre', '', 'Notifications closes', 'UI', 'P3')
add(M, 'Quick-access hub', P, 'All links inside hub panels resolve', PRE_HOME, '1. Open each hub\n2. Visit every link', '', 'All 200', 'Functional', 'P1', IN)
add(M, 'Announcements', P, 'Ticker shows counter (e.g. 1/3)', PRE_HOME, '1. Look at announcement strip', '', 'Counter and one announcement visible', 'UI', 'P2', HP)
add(M, 'Announcements', P, 'Next arrow moves to next announcement', PRE_HOME, '1. Click Next', '', 'Counter increments', 'Functional', 'P2', HP)
add(M, 'Announcements', P, 'Previous arrow moves back', PRE_HOME, '1. Click Next\n2. Click Previous', '', 'Counter decrements. SUSPECTED DEFECT: Previous does nothing', 'Functional', 'P2', HP)
add(M, 'Announcements', P, 'Pause stops auto-rotation; Play resumes', PRE_HOME, '1. Click Pause\n2. Wait 10s\n3. Click Play', '', 'No change while paused; rotates after Play; label toggles Pause/Play', 'Functional', 'P2', HP)
add(M, 'Announcements', P, 'EDGE: Next on last / Previous on first', PRE_HOME, '1. Go to 3/3, click Next\n2. Go to 1/3, click Previous', '', 'Wraps around (or button disabled) — no blank slide', 'Edge', 'P3')
add(M, 'Announcements', P, 'Announcement links open valid pages', PRE_HOME, '1. Click each announcement', '', 'All 200', 'Functional', 'P1', HP)
add(M, 'Announcements', P, 'Rotation pauses on hover/focus (WCAG 2.2.2)', PRE_HOME, '1. Hover / Tab into ticker', '', 'Auto-rotation stops while hovered/focused', 'Accessibility', 'P3')
add(M, 'Current Rates', P, 'Policy Rates tab shows Repo Rate %', PRE_HOME, '1. View Current Rates', '', 'Policy Repo Rate shows a valid % (0–20), e.g. 5.25%', 'Functional', 'P1', HP)
for tab in ['Policy Rates', 'Reserve Ratios', 'Exchange Rates', 'Lending/ Deposit Rates', 'Market Trends']:
    add(M, 'Current Rates', P, f'"{tab}" tab shows values', PRE_HOME, f'1. Click "{tab}" tab', '', 'Tab selected; panel shows numeric values; no NaN/undefined/blank', 'Functional', 'P1', HP)
add(M, 'Current Rates', P, 'Rates match official source', PRE_HOME, '1. Compare each value with RBI data release', 'rbi.org.in current rates', 'Values and "as on" dates match', 'Content', 'P1')
add(M, 'Current Rates', P, 'Rate tabs keyboard navigable (arrow keys)', PRE_HOME, '1. Tab into rate tabs\n2. Use ← →', '', 'Focus moves between tabs; Enter/Space activates', 'Accessibility', 'P3')
add(M, "Governor's Desk", P, "Governor's Desk section renders", PRE_HOME, '1. Scroll to Governor\'s Desk', '', 'Photo, name, links visible', 'UI', 'P2', HP)
add(M, "Governor's Desk", P, "Governor's Desk links resolve", PRE_HOME, '1. Click each link', '', 'All 200. KNOWN DEFECT: /governor-monetary-policy-statements → 404', 'Functional', 'P1', LK)
add(M, 'Latest Updates', P, "What's New / Sections Updated Today tabs switch", PRE_HOME, '1. Click each tab', '', 'Tab selected; its panel shows', 'Functional', 'P1', HP)
add(M, 'Latest Updates', P, "What's New items show type, date and title", PRE_HOME, "1. View What's New", '', 'Each item: category (e.g. Press Releases) | date (Month DD, YYYY) | title', 'UI', 'P2', HP)
add(M, 'Latest Updates', P, "What's New items open detail pages", PRE_HOME, '1. Click an item', '', 'Correct detail page opens', 'Functional', 'P1', HP)
add(M, 'Latest Updates', P, 'Items sorted newest first', PRE_HOME, "1. Read dates in What's New", '', 'Dates in descending order', 'Functional', 'P2')
add(M, 'Latest Updates', P, 'EDGE: "Sections Updated Today" with nothing updated', PRE_HOME, '1. Check on a day with no updates', '', 'Friendly empty message, not blank panel', 'Edge', 'P3')

# ============================== 7. LISTINGS ==============================
M = 'Listing Pages'
listings = ['/press-releases', '/notifications/master-directions', '/notifications/master-circulars', '/notifications/amendment-directions', '/notifications/standalone-circulars', '/notifications/circular-withdrawn', '/notifications/draft-notifications', '/notifications/draft-directions-re-wise', '/notifications/index-to-rbi-circulars', '/faqs', '/speech-and-media/media-interactions', '/speech-and-media/podcasts', '/publications/working-papers']
for L in listings:
    add(M, 'Load', L, 'Listing loads with title, breadcrumb and items', PRE_ANY, f'1. Open {L}', '', 'H1, breadcrumb (Home > …), item cards with date + title', 'Functional', 'P1', LS)
    add(M, 'Detail link', L, 'First item opens its detail page', PRE_ANY, f'1. Open {L}\n2. Click first item', '', 'Detail page title matches card title', 'Functional', 'P1', LS)
M = 'Listing Filters'
add(M, 'Year', 'All listings', 'Year dropdown lists years and filters', PRE_ANY, '1. Click Select year\n2. Choose 2025', '2025', 'Button shows 2025; only 2025 items', 'Functional', 'P1', TP)
add(M, 'Month', 'All listings', 'Month dropdown lists 12 months and filters', PRE_ANY, '1. Click Select month\n2. Choose March', 'March', 'Only March items', 'Functional', 'P1', TP)
add(M, 'Year+Month', 'All listings', 'Year + Month together', PRE_ANY, '1. Year 2024\n2. Month June', '2024 / June', 'Only June 2024 items', 'Functional', 'P2')
add(M, 'Dropdown', 'All listings', 'Dropdowns keyboard operable', PRE_ANY, '1. Tab to year\n2. Enter, ↓, Enter, Esc', '', 'Opens, moves, selects, closes', 'Accessibility', 'P2', TP)
add(M, 'Date range', 'All listings', 'Custom date range filters', PRE_ANY, '1. Click "Filter using custom date"\n2. Pick From and To\n3. Apply', '01-Jan-2024 to 31-Mar-2024', 'Only items in range', 'Functional', 'P1', TP)
add(M, 'Date range', 'All listings', 'EDGE: From date after To date', PRE_ANY, '1. From 31-Dec-2026, To 01-Jan-2020\n2. Apply', '', 'Validation message or dates swapped; no error', 'Edge', 'P1', TP)
add(M, 'Date range', 'All listings', 'EDGE: future date range', PRE_ANY, '1. Range in the future', '01-Jan-2030 to 31-Dec-2030', 'Future dates disabled or no-results message', 'Edge', 'P2')
add(M, 'Date range', 'All listings', 'EDGE: same From and To date', PRE_ANY, '1. From = To', '03-Jul-2026', 'Items of that day only', 'Edge', 'P3')
add(M, 'Date range', 'All listings', 'EDGE: only From date set', PRE_ANY, '1. Set From only\n2. Apply', '', 'Validation or open-ended range', 'Edge', 'P3')
add(M, 'Date range', 'All listings', 'Clear resets date range', PRE_ANY, '1. Set range\n2. Clear', '', 'Dates back to "Select date"; full list', 'Functional', 'P2', TP)
add(M, 'Function filter', '/press-releases, /faqs', 'Category checkbox filters and shows count', PRE_ANY, '1. Tick "About Us (137)"', '', 'Results = 137; list only that category', 'Functional', 'P1', LS)
add(M, 'Function filter', '/press-releases, /faqs', 'Multiple categories combine (OR)', PRE_ANY, '1. Tick two categories', '', 'Results = sum of both counts (or documented logic)', 'Functional', 'P2')
add(M, 'Function filter', '/press-releases, /faqs', 'Category search box filters the checkbox list', PRE_ANY, '1. Type "Co-operative" in category search', 'Co-operative', 'Only matching categories shown', 'Functional', 'P2', LS)
add(M, 'Function filter', '/faqs', '"View N more" expands category list', PRE_ANY, '1. Click "View 9 more"', '', 'More categories shown; button toggles to "View less"', 'Functional', 'P3', TP)
add(M, 'Function filter', 'All listings', 'EDGE: category search with no match', PRE_ANY, '1. Type "zzzz"', 'zzzz', '"No matching categories" message', 'Edge', 'P3')
add(M, 'Keyword', '/press-releases', 'Keyword search inside listing', PRE_ANY, '1. Type "Treasury Bills" in listing search\n2. Enter', 'Treasury Bills', 'Only matching items', 'Functional', 'P1', LS)
add(M, 'Sort', '/press-releases', 'Sort Newest first / Oldest first', PRE_ANY, '1. Change sort to Oldest first', '', 'Order reverses; dates ascending', 'Functional', 'P2', LS)
add(M, 'Page size', 'All listings', 'Items per page (10/20/50)', PRE_ANY, '1. Change to 50', '', '50 cards shown', 'Functional', 'P2', LS)
add(M, 'Pagination', 'All listings', 'Page numbers, Next, Previous work', PRE_ANY, '1. Click 2, Next, Previous', '', 'Items change; current page highlighted; Previous disabled on page 1', 'Functional', 'P1', LS)
add(M, 'Pagination', 'All listings', 'Last page reachable', PRE_ANY, '1. Click last page number (e.g. 3074)', '', 'Last items load; Next disabled', 'Functional', 'P2')
add(M, 'Pagination', 'All listings', 'EDGE: page number far beyond last in URL', PRE_ANY, '1. Open ?page=999999', '', 'Empty state or last page; no 500', 'Edge', 'P2', TP)
add(M, 'Pagination', 'All listings', 'EDGE: garbage query params', PRE_ANY, '1. Open ?year=abcd&month=99&page=-1', '', 'Page loads with defaults; no 500', 'Edge', 'P2', TP)
add(M, 'Filters', 'All listings', 'Filters kept in URL / after Back', PRE_ANY, '1. Apply filters\n2. Open item\n3. Back', '', 'Same filters and page restored', 'Edge', 'P2')
add(M, 'Filters', 'All listings', '"Show N results" matches items shown', PRE_ANY, '1. Apply filter\n2. Compare count and cards', '', 'Counts equal', 'Functional', 'P2', TP)
add(M, 'Filters', 'All listings', 'Filter with zero results shows empty state', PRE_ANY, '1. Choose filters with no data', '', '"No results" message + Clear filters', 'Edge', 'P2')
add(M, 'Mobile', 'All listings', 'Mobile filter panel opens/closes', 'Mobile 375px', '1. Tap Filters\n2. Tap Close filters', '', 'Panel slides in/out; body not scrollable behind', 'UI', 'P2', TP)

# ============================== 8. DETAIL PAGES ==============================
M = 'Detail Pages'
for kind, list_url in [('Press release', '/press-releases'), ('Notification / circular', '/notifications/master-directions'), ('FAQ', '/faqs'), ('Media interaction transcript', '/speech-and-media/media-interactions'), ('Publication', '/publications/annual')]:
    add(M, kind, list_url + '/…', f'{kind}: title, date, breadcrumb present', PRE_ANY, f'1. Open any {kind.lower()}', '', 'H1, date (<time>), breadcrumb, Last Updated On', 'Functional', 'P1', LS if kind == 'Press release' else PG)
    add(M, kind, list_url + '/…', f'{kind}: PDF link opens a valid PDF in new tab', PRE_ANY, '1. Click PDF link', '', 'Opens in new tab; file is a real PDF; size label (e.g. 209 KB) correct', 'Functional', 'P1', LS if kind == 'Press release' else LK)
add(M, 'Breadcrumb', 'Detail pages', 'Breadcrumb links go back to listing / home', PRE_ANY, '1. Click each breadcrumb', '', 'Correct pages open', 'Functional', 'P2', LS)
add(M, 'Content', 'Detail pages', 'Tables inside content are readable and scroll on mobile', PRE_ANY, '1. Open a circular with a table\n2. Check at 375px', '', 'Table scrolls horizontally inside its box; page does not', 'UI', 'P2', RS)
add(M, 'Content', 'Detail pages', 'Long titles wrap without overflow', PRE_ANY, '1. Open an item with a very long title', '', 'Title wraps; no overflow/cut-off', 'Edge', 'P3', PG)
add(M, 'Print', 'Detail pages', 'Print view is clean', PRE_ANY, '1. Ctrl+P', '', 'Header/menus hidden; content fits A4', 'UI', 'P3')
add(M, 'Edge', 'Detail pages', 'EDGE: invalid/removed item slug', PRE_ANY, '1. Open /press-releases/does-not-exist-123', '', 'HTTP 404 + friendly page', 'Edge', 'P2', LS)
add(M, 'Edge', 'Detail pages', 'EDGE: slug with trailing slash / uppercase', PRE_ANY, '1. Open the same item URL with "/" and in UPPERCASE', '', 'Redirects to canonical URL or loads same page', 'Edge', 'P3')

# ============================== 9. STATIC / TEMPLATE PAGES ==============================
M = 'Static & Special Pages'
for p in ['/banker-to-banks/overview', '/enforcement-department/overview', '/issuer-of-currency/overview', '/monetary-policy/overview', '/payment-and-settlement-systems/overview', '/citizens-charter', '/right-to-information-act']:
    add(M, 'On-this-page index', p, 'Index Hide/Show toggle and in-page search', PRE_ANY, f'1. Open {p}\n2. Click Hide / Show\n3. Search a word in "Search this page"', 'a word from a heading', 'Index collapses/expands; search jumps to/highlights match', 'Functional', 'P2', TP)
add(M, 'On-this-page index', 'Static pages', 'Index links jump to their section', PRE_ANY, '1. Click an index link', '', 'Scrolls to section; heading not hidden under sticky header', 'Functional', 'P2', TP)
add(M, 'On-this-page index', 'Static pages', 'EDGE: in-page search with no match', PRE_ANY, '1. Search "zzqqxx"', '', '"No matches" message; no error', 'Edge', 'P3', TP)
add(M, 'About Us', '/about-us', 'Building carousel next/prev/dots/pause', PRE_ANY, '1. Pause\n2. Next, Previous, dot 3', '', 'Slides change accordingly', 'Functional', 'P2', TP)
add(M, 'About Us', '/about-us', 'EDGE: carousel wraps from last to first', PRE_ANY, '1. Go to last slide\n2. Next', '', 'Goes to first slide', 'Edge', 'P3', TP)
add(M, 'About Us', '/about-us', '"More sections" rail expands', PRE_ANY, '1. Click More sections', '', 'Extra sections listed', 'Functional', 'P3', TP)
add(M, 'Organisation', '/about-us/organisation', 'Each org-chart node expands/collapses', PRE_ANY, '1. Click each Executive Director node', '', 'Details show/hide; aria-expanded toggles', 'Functional', 'P2', TP)
add(M, 'Central Board', '/about-us/central-board', 'Board members table correct', PRE_ANY, '1. Review table', '', 'Names/designations correct; table has headers', 'Content', 'P2')
add(M, 'Offices', '/about-us/offices', 'Map View / List View tabs', PRE_ANY, '1. Click List View\n2. Click Map View', '', 'Views switch', 'Functional', 'P1', TP)
add(M, 'Offices', '/about-us/offices', 'Map pin shows office details', PRE_ANY, '1. Click Srinagar pin', '', 'Office name + address shown', 'Functional', 'P1', TP)
add(M, 'Offices', '/about-us/offices', 'All 34 map pins clickable, no overlap', PRE_ANY, '1. Click every pin', '', 'Each shows its office; no pin hidden under another', 'Functional', 'P2', TP)
add(M, 'Offices', '/about-us/offices', 'Office search by city', PRE_ANY, '1. Search Mumbai', 'Mumbai', 'Mumbai office(s) shown', 'Functional', 'P1', TP)
add(M, 'Offices', '/about-us/offices', 'EDGE: office search no match', PRE_ANY, '1. Search Atlantis', 'Atlantis', '"No offices found"', 'Edge', 'P2', TP)
add(M, 'Offices', '/about-us/offices', 'EDGE: office search case/space insensitive', PRE_ANY, '1. Search "  mUmBaI  "', '', 'Same as "Mumbai"', 'Edge', 'P3', TP)
add(M, 'Offices', '/about-us/offices', 'Office phone/email links work', PRE_ANY, '1. Click tel:/mailto: links', '', 'Dialer / mail client opens with correct value', 'Functional', 'P3')
add(M, 'Departments', '/about-us/departments', 'Department search filters list', PRE_ANY, '1. Search "Currency"', 'Currency', 'Only matching departments', 'Functional', 'P2', TP)
add(M, 'Departments', '/about-us/departments', 'EDGE: clearing search restores list', PRE_ANY, '1. Search\n2. Clear box', '', 'Full list back', 'Edge', 'P3', TP)
add(M, 'RBI History', '/about-us/rbi-history', 'Photo gallery next/dots', PRE_ANY, '1. Click Next photograph, dots', '', 'Photos change; captions update', 'Functional', 'P3', TP)
add(M, 'RBI History', '/about-us/rbi-history', 'Volumes carousel + Read More', PRE_ANY, '1. Next Volumes\n2. Read More on each', '', 'Text expands; aria-expanded=true', 'Functional', 'P3', TP)
add(M, 'Legal framework', '/legal-framework/act', 'Act links open India Code in new tab', PRE_ANY, '1. Click each Act card', '', 'Opens indiacode.nic.in in new tab with rel=noopener', 'Functional', 'P2', LK)
add(M, 'Publications', '/publications/*', 'Publication category pages list cards', PRE_ANY, '1. Open annual, monthly, weekly… pages', '', 'Cards link to publication pages', 'Functional', 'P2', PG)
add(M, 'Complaints', '/complaints', 'Grievance redress table + links', PRE_ANY, '1. Open /complaints', '', 'Table readable; CMS / Ombudsman links work', 'Functional', 'P1', PG)
add(M, 'Soft 404', '/speech-and-media', 'Speeches & Media landing page exists', PRE_ANY, '1. Open /speech-and-media', '', 'Real landing page. KNOWN DEFECT: shows "not found" with HTTP 200', 'Functional', 'P1', TP)
add(M, 'Podcasts', '/speech-and-media/podcasts', 'Podcast audio plays', PRE_ANY, '1. Open a podcast\n2. Play', '', 'Audio plays; pause/seek/volume work; transcript if any', 'Functional', 'P2')

# ============================== 10. FOOTER ==============================
M = 'Footer'
add(M, 'Links', 'All pages', 'Quick Links all resolve', PRE_ANY, '1. Click every footer link', '', 'All 200. KNOWN DEFECTS: Opportunities, Important Websites, Bank Holidays, RBI Clarifications, Forms, RBI\'s Vision and Values → 404', 'Functional', 'P1', LK)
add(M, 'Social', 'All pages', 'Social icons open correct RBI profiles in new tab', PRE_ANY, '1. Click YouTube, X, Instagram, Facebook, LinkedIn', '', 'Official RBI accounts; new tab; rel=noopener', 'Functional', 'P2', HP)
add(M, 'Need Help', 'All pages', 'Need Help section links work', PRE_ANY, '1. Click each Need Help link', '', 'Correct pages / contact info', 'Functional', 'P2', LK)
add(M, 'Content', 'All pages', 'Copyright year and "Last updated" correct', PRE_ANY, '1. Read footer text', '', 'Current year; correct dates', 'Content', 'P3')
add(M, 'Layout', 'All pages', 'Footer layout on mobile', 'Mobile 375px', '1. Scroll to footer', '', 'Columns stack; links tappable (≥24px)', 'UI', 'P2', RS)

# ============================== 11. ALL PAGES (automated sweeps) ==============================
M = 'Every Page (automated sweep)'
add(M, 'Health', 'Every crawled page', 'Page returns HTTP < 400', 'npm run crawl done', 'Automated: open every crawled URL', '', 'Status 200 (or redirect)', 'Functional', 'P1', PG)
add(M, 'Health', 'Every crawled page', 'No uncaught JavaScript errors', 'npm run crawl done', 'Automated', '', 'No page errors', 'Functional', 'P1', PG)
add(M, 'Health', 'Every crawled page', 'No console errors / failed requests', 'npm run crawl done', 'Automated', '', 'No console errors; no 4xx/5xx sub-requests', 'Functional', 'P2', PG)
add(M, 'Health', 'Every crawled page', 'No broken images; all images have alt', 'npm run crawl done', 'Automated (scrolls to load lazy images)', '', 'All images load; alt present', 'UI', 'P2', PG)
add(M, 'Health', 'Every crawled page', 'Exactly one visible H1', 'npm run crawl done', 'Automated', '', '1 H1', 'SEO', 'P3', PG)
add(M, 'Health', 'Every crawled page', 'No template leaks (undefined, NaN, lorem ipsum)', 'npm run crawl done', 'Automated', '', 'None found', 'Content', 'P2', PG)
add(M, 'Health', 'Every crawled page', 'No horizontal scroll at desktop', 'npm run crawl done', 'Automated', '', 'No overflow', 'UI', 'P2', PG)
add(M, 'Links', 'Every crawled page', 'Every link (internal, PDF, external) resolves', 'npm run crawl done', 'Automated link checker', '', 'All < 400 (external 403/429 = warning)', 'Functional', 'P1', LK)
add(M, 'Buttons', 'Every crawled page', 'Every button/tab/accordion/dropdown works', 'npm run crawl done', 'Automated auto-click tester', '', 'No JS error, no 404 after click, not covered; dead buttons reported', 'Functional', 'P1', IN)
add(M, 'Inputs', 'Every crawled page', 'Every input handles edge-case data', 'npm run crawl done', 'Automated: empty, spaces, 500 chars, special chars, XSS, SQL-like, Unicode, numbers', '', 'No script execution, no JS error, no 5xx', 'Security', 'P1', IN)
add(M, 'Accessibility', 'Every crawled page', 'axe-core WCAG 2.1 AA scan', 'npm run crawl done', 'Automated', '', 'No serious/critical violations', 'Accessibility', 'P1', AX)
add(M, 'Locators', 'Key pages', 'All XPath locators still match after deploy', '', 'Automated', '', 'Every XPath matches ≥1 element', 'Functional', 'P1', XL)

# ============================== 12. RESPONSIVE & COMPATIBILITY ==============================
M = 'Responsive & Compatibility'
for w, dev in [(375, 'iPhone SE / small phone'), (390, 'iPhone 14'), (768, 'iPad portrait'), (1024, 'iPad landscape / small laptop'), (1280, 'Laptop'), (1440, 'Desktop'), (1920, 'Full HD')]:
    add(M, 'Breakpoints', 'Key pages', f'Layout at {w}px ({dev})', f'Viewport {w}px', '1. Open home, listing, detail, offices, search', '', 'No overflow, overlaps or cut text; menus usable', 'UI', 'P1' if w in (375, 1440) else 'P2', RS if w in (375, 768, 1024, 1280) else '')
add(M, 'Orientation', 'Key pages', 'Rotate phone portrait ↔ landscape', 'Mobile', '1. Rotate device on each key page', '', 'Layout adapts; no lost state', 'UI', 'P3')
add(M, 'Zoom', 'All pages', 'Browser zoom 200% (WCAG 1.4.4) and 400% reflow (1.4.10)', PRE_ANY, '1. Ctrl + to 200% / 400%', '', 'Content readable; no loss of function; single-column reflow at 400%', 'Accessibility', 'P2')
add(M, 'Touch', 'All pages', 'Tap targets ≥ 24×24px', 'Mobile', '1. Check header, footer, pager, filters', '', 'All targets large enough', 'Accessibility', 'P2', RS)
for b in ['Chrome (latest)', 'Edge (latest)', 'Firefox (latest)', 'Safari macOS (latest)', 'Safari iOS', 'Chrome Android', 'Samsung Internet']:
    add(M, 'Browsers', 'Key pages', f'Smoke test on {b}', PRE_ANY, '1. Home, menu, search, listing filters, detail PDF, a11y toolbar', '', 'Same behaviour and look as Chrome', 'Compatibility', 'P1' if 'Chrome' in b or 'Safari iOS' in b else 'P2')

# ============================== 13. ACCESSIBILITY (manual) ==============================
M = 'Accessibility (manual)'
add(M, 'Keyboard', 'All pages', 'Whole site usable with keyboard only', PRE_ANY, '1. Unplug mouse\n2. Navigate menus, search, filters, carousels, toolbar', '', 'Everything reachable; logical order; no traps', 'Accessibility', 'P1')
add(M, 'Focus', 'All pages', 'Visible focus indicator on every control', PRE_ANY, '1. Tab through page', '', 'Clear focus ring everywhere', 'Accessibility', 'P1')
add(M, 'Screen reader', 'Key pages', 'NVDA / VoiceOver read page correctly', 'NVDA or VoiceOver on', '1. Read home, listing, detail with screen reader', '', 'Landmarks, headings, button names, tab states, expanded states announced', 'Accessibility', 'P1')
add(M, 'Contrast', 'All pages', 'Text contrast ≥ 4.5:1 (light and dark theme)', PRE_ANY, '1. Check with contrast tool', '', 'All text passes AA', 'Accessibility', 'P2', AX)
add(M, 'Motion', 'All pages', 'Respects "reduce motion" OS setting', 'OS reduce motion on', '1. Load home, about-us', '', 'Carousels do not auto-animate', 'Accessibility', 'P3')
add(M, 'Forms', 'All inputs', 'Every input has a label', PRE_ANY, '1. Inspect search, filter, date inputs', '', 'Label or aria-label present', 'Accessibility', 'P2', AX)
add(M, 'Documents', 'PDF links', 'PDF links say "opens in new tab" + size', PRE_ANY, '1. Inspect PDF links', '', 'Text/aria indicates PDF, new tab and size', 'Accessibility', 'P3')

# ============================== 14. PERFORMANCE ==============================
M = 'Performance'
add(M, 'Load time', 'Key pages', 'Home LCP < 2.5 s, CLS < 0.1, INP < 200 ms', PRE_ANY, '1. Lighthouse / PageSpeed on home (mobile + desktop)', '', 'Core Web Vitals pass', 'Performance', 'P1')
add(M, 'Load time', 'Listing pages', 'Listing pages respond < 2 s', PRE_ANY, '1. Measure TTFB on notifications and speech pages', '', '< 2 s. OBSERVED: 3–4 s on several notifications/speech pages', 'Performance', 'P2', PG)
add(M, 'Weight', 'Home', 'Homepage HTML size reasonable', PRE_ANY, '1. Check document size', '', 'HTML < 300 KB. OBSERVED: ~430 KB (inlined SSR state)', 'Performance', 'P3')
add(M, 'Images', 'All pages', 'Images optimised (WebP/AVIF, lazy, sized)', PRE_ANY, '1. Lighthouse image audits', '', 'No oversized images; width/height set', 'Performance', 'P2')
add(M, 'Caching', 'All pages', 'Static assets cached', PRE_ANY, '1. Check Cache-Control on JS/CSS/images', '', 'Long max-age with hashed filenames', 'Performance', 'P3')
add(M, 'Load', 'Key pages', 'Site handles expected concurrent users', 'Load-test tool (k6/JMeter)', '1. Ramp 200 users on home + search + listings', '', 'p95 < 3 s; no 5xx', 'Performance', 'P2')
add(M, 'Network', 'Key pages', 'Usable on slow 3G', 'DevTools throttling Slow 3G', '1. Load home and a listing', '', 'Content appears progressively; no broken layout', 'Performance', 'P3')

# ============================== 15. SEO ==============================
M = 'SEO'
add(M, 'Files', '/sitemap.xml', 'sitemap.xml exists and lists pages', PRE_ANY, '1. Open /sitemap.xml', '', 'Valid XML sitemap. KNOWN DEFECT: 404', 'SEO', 'P2', LS)
add(M, 'Files', '/robots.txt', 'robots.txt exists', PRE_ANY, '1. Open /robots.txt', '', 'Valid robots.txt pointing to sitemap. KNOWN DEFECT: 404', 'SEO', 'P2', LS)
add(M, 'Meta', 'Every page', 'Unique title and meta description per page', PRE_ANY, 'Automated + manual review', '', 'Present and unique', 'SEO', 'P2', PG)
add(M, 'Meta', 'Every page', 'Canonical URL present', PRE_ANY, '1. Inspect <link rel=canonical>', '', 'Points to production URL of the page', 'SEO', 'P3')
add(M, 'Meta', 'Every page', 'Open Graph / Twitter tags for sharing', PRE_ANY, '1. Inspect og:title, og:image', '', 'Present; share preview correct', 'SEO', 'P3')
add(M, 'Errors', 'Unknown URL', 'Unknown URL returns real 404', PRE_ANY, '1. Open /this-page-does-not-exist', '', 'HTTP 404 + friendly page with Home link', 'SEO', 'P1', LS)

# ============================== 16. ERROR HANDLING & RESILIENCE ==============================
M = 'Error Handling'
add(M, '404', 'Unknown URL', '404 page shows header, search and Home link', PRE_ANY, '1. Open random URL', '', 'Branded 404 with navigation', 'UI', 'P2', LS)
add(M, 'Offline', 'Any page', 'EDGE: network drops while filtering', PRE_ANY, '1. DevTools → Offline\n2. Change a filter', '', 'Friendly error / retry, not blank page', 'Edge', 'P3')
add(M, 'JS disabled', 'Key pages', 'EDGE: JavaScript disabled (SSR fallback)', 'JS disabled in browser', '1. Load home and a listing', '', 'Content readable via server render; links work', 'Edge', 'P3')
add(M, 'Double click', 'Forms / filters', 'EDGE: double-click Apply / Search', PRE_ANY, '1. Double-click Apply quickly', '', 'Single request; no duplicate results', 'Edge', 'P3')
add(M, 'Back/Forward', 'All pages', 'EDGE: browser Back/Forward through SPA navigation', PRE_ANY, '1. Home → listing → detail → Back → Back → Forward', '', 'Correct pages & scroll positions', 'Edge', 'P2')
add(M, 'Refresh', 'All pages', 'EDGE: refresh on every page type keeps state', PRE_ANY, '1. Apply filters\n2. F5', '', 'Same page and filters (if in URL)', 'Edge', 'P3')
add(M, 'Deep link', 'Any page', 'EDGE: open deep link in new tab / incognito', PRE_ANY, '1. Copy a detail URL\n2. Open in new incognito window (after login)', '', 'Page loads directly', 'Edge', 'P2')

# ----------------------------------------------------------------------------------------------
defects = [
    ('D-01', 'Broken link', 'Header → Functions menu', '"Regulation" link goes to /regulation → 404', 'All pages', 'High', '/regulation'),
    ('D-02', 'Broken link', 'Footer', '"Opportunities" → 404', 'All pages', 'High', '/opportunities'),
    ('D-03', 'Broken link', 'Footer', '"Important Websites" → 404', 'All pages', 'High', '/important-websites'),
    ('D-04', 'Broken link', 'Footer', '"Bank Holidays" → 404', 'All pages', 'High', '/bank-holidays'),
    ('D-05', 'Broken link', 'Footer', '"RBI Clarifications" → 404', 'All pages', 'High', '/rbi-clarifications'),
    ('D-06', 'Broken link', 'Footer', '"Forms" → 404', 'All pages', 'High', '/forms'),
    ('D-07', 'Broken link', 'Footer', '"RBI\'s Vision and Values" → 404', 'All pages', 'High', '/about-rbi/vision-and-values'),
    ('D-08', 'Broken link', 'Site-wide', '"RBI Museum" → 404', 'All pages', 'Medium', '/about-us/rbi-museum'),
    ('D-09', 'Broken link', "Homepage – Governor's Desk", '"Monetary Policy Statements, Governor…" → 404', '/', 'Medium', '/governor-monetary-policy-statements'),
    ('D-10', 'Soft 404', 'Speeches & Media landing', 'Returns HTTP 200 but renders the "not found" page', '/speech-and-media', 'High', '/speech-and-media'),
    ('D-11', 'Feature missing', 'Header – Change language', 'Button has no action; Hindi site /hi returns 404', 'All pages', 'High', '/hi'),
    ('D-12', 'Functional', 'Homepage – announcement ticker', '"Previous announcement" does not move back (Next works) — confirm manually', '/', 'Medium', '/'),
    ('D-13', 'SEO', 'Site', '/sitemap.xml returns 404', '-', 'Medium', '/sitemap.xml'),
    ('D-14', 'SEO', 'Site', '/robots.txt returns 404', '-', 'Low', '/robots.txt'),
    ('D-15', 'Performance', 'Notifications & Speech listings', 'HTML response 3–4 s under light concurrency', 'notifications/*, speech-and-media/*', 'Medium', '/notifications/draft-directions-re-wise'),
]

inventory_raw = """/=TA /about-us=TD /about-us/organisation=AD /about-us/central-board=B /about-us/departments= /about-us/offices=T /about-us/rbi-history=A /about-us/communication-policy=ID /about-us/sources-of-information=IB /about-us/media-kit= /press-releases=FPACD /banker-to-banks/overview=I /consumer-education-and-protection-department/overview=I /issuer-of-currency/overview=I /banker-and-debt-manager-to-government/overview=I /enforcement-department/overview=ID /external-investments-and-operations/overview= /fintech/overview=A /financial-inclusion-and-development/overview=ID /financial-markets/overview=ID /financial-stability-unit/overview=I /foreign-exchange-management/overview=I /functions/international-relations= /monetary-policy/overview=IBD /payment-and-settlement-systems/overview=IBD /regulation=404 /research-and-data/overview=I /functions/supervision= /notifications/master-directions=FPACD /notifications/master-circulars=FPACD /notifications/amendment-directions=FPACD /notifications/index-to-rbi-circulars=FPAB /notifications/standalone-circulars=FPACD /notifications/circular-withdrawn=FPACD /notifications/draft-notifications=FPACD /notifications/draft-directions-re-wise=FPAC /speech-and-media=X /speech-and-media/media-interactions=FPA /speech-and-media/podcasts=FA /speech-and-media/memorial-lectures= /legal-framework/act=CD /legal-framework/rules=CD /legal-framework/regulations=CD /legal-framework/schemes=CD /issuer-of-currency/mani= /complaints=BD /citizens-charter=IBD /regulated-entities=D /faqs=FPAC /right-to-information-act=IBD /publications/biennial=C /publications/annual=C /publications/half-yearly= /publications/quarterly=C /publications/bi-monthly=C /publications/monthly=C /publications/weekly=C /publications/occasional=CD /publications/reports=C /publications/working-papers=FPAD /external-research-schemes= /statistics/public-debt-statistics="""
FEAT = {'F': 'Filters', 'P': 'Pagination', 'T': 'Tabs', 'A': 'Accordions/toggles', 'B': 'Tables', 'C': 'Cards', 'D': 'PDF links', 'I': 'On-page index + search', 'X': 'SOFT 404', '4': 'HTTP 404'}
def template_of(path, code):
    if path == '/': return 'Homepage'
    if code == '404': return 'Broken (404)'
    if 'X' in code: return 'Not-found template (soft 404)'
    if 'F' in code: return 'Listing with filters'
    if 'I' in code: return 'Static page with index'
    if 'C' in code: return 'Card grid'
    return 'Static page'
inventory = []
for tok in inventory_raw.split():
    path, code = tok.split('=')
    feats = 'HTTP 404' if code == '404' else ', '.join(FEAT[c] for c in code if c in FEAT)
    inventory.append((path, template_of(path, code), feats or '—'))

# ============================== WORKBOOK ==============================
wb = Workbook()
ARIAL = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
body_font = Font(name=ARIAL, size=10)
bold = Font(name=ARIAL, bold=True, size=10)
title_font = Font(name=ARIAL, bold=True, size=14, color='1F3864')
input_fill = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top', wrap_text=True)

# ---- Test Cases sheet
ws = wb.active
ws.title = 'Test Cases'
cols = ['TC ID', 'Module', 'Feature', 'Page / URL', 'Test Scenario', 'Preconditions', 'Test Steps', 'Test Data', 'Expected Result', 'Type', 'Priority', 'Automation', 'Automated In (spec)', 'Status', 'Actual Result', 'Defect ID / Remarks', 'Tested By', 'Test Date']
widths = [10, 20, 18, 26, 42, 20, 42, 20, 46, 13, 8, 12, 30, 11, 30, 24, 12, 11]
ws.append(cols)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(row=1, column=i)
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
ws.row_dimensions[1].height = 30

prefix = {}
counter = {}
for m in dict.fromkeys(c[0] for c in cases):
    words = [w for w in m.replace('&', ' ').replace('(', ' ').replace(')', ' ').split() if w]
    p = ''.join(w[0] for w in words[:3]).upper()
    while p in prefix.values():
        p += words[0][1].upper()
    prefix[m] = p

for idx, (module, feature, page, scenario, pre, steps, data, expected, typ, pri, auto) in enumerate(cases, start=2):
    counter[module] = counter.get(module, 0) + 1
    tcid = f'TC-{prefix[module]}-{counter[module]:03d}'
    url = page if not page.startswith('/') or '…' in page or ',' in page or '*' in page else page
    row = [tcid, module, feature, url, scenario, pre, steps, data, expected, typ, pri, 'Automated' if auto else 'Manual', auto, 'Not Run', '', '', '', '']
    ws.append(row)
    for ci in range(1, len(cols) + 1):
        c = ws.cell(row=idx, column=ci)
        c.font = body_font
        c.alignment = center if ci in (1, 10, 11, 12, 14, 18) else wrap
        c.border = border
        if ci in (14, 15, 16, 17, 18):
            c.fill = input_fill
last = ws.max_row
ws.freeze_panes = 'F2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'

dv = DataValidation(type='list', formula1='"Not Run,Pass,Fail,Blocked,N/A"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'N2:N{last}')
for status, color in [('Pass', 'C6EFCE'), ('Fail', 'FFC7CE'), ('Blocked', 'FFEB9C'), ('N/A', 'D9D9D9')]:
    ws.conditional_formatting.add(f'N2:N{last}', CellIsRule(operator='equal', formula=[f'"{status}"'], fill=PatternFill('solid', fgColor=color)))
ws.conditional_formatting.add(f'K2:K{last}', CellIsRule(operator='equal', formula=['"P1"'], font=Font(name=ARIAL, bold=True, color='C00000')))
ws.conditional_formatting.add(f'I2:I{last}', FormulaRule(formula=[f'ISNUMBER(SEARCH("KNOWN DEFECT",I2))'], font=Font(name=ARIAL, color='C00000')))

# ---- Summary sheet
sm = wb.create_sheet('Summary', 0)
sm['A1'] = 'RBI Website (stg-rbi.webc.in) — Test Case Summary'
sm['A1'].font = title_font
sm['A2'] = 'All figures below are formulas over the "Test Cases" sheet — they update as you set Status.'
sm['A2'].font = Font(name=ARIAL, italic=True, size=9, color='595959')
sm['A4'], sm['B4'] = 'Total test cases', f"=COUNTA('Test Cases'!A2:A{last})"
sm['A5'], sm['B5'] = 'Automated', f"=COUNTIF('Test Cases'!L2:L{last},\"Automated\")"
sm['A6'], sm['B6'] = 'Manual', f"=COUNTIF('Test Cases'!L2:L{last},\"Manual\")"
sm['A7'], sm['B7'] = 'Automation coverage', '=IF(B4=0,0,B5/B4)'
sm['A8'], sm['B8'] = 'Executed (Pass + Fail + Blocked)', f"=COUNTIF('Test Cases'!N2:N{last},\"Pass\")+COUNTIF('Test Cases'!N2:N{last},\"Fail\")+COUNTIF('Test Cases'!N2:N{last},\"Blocked\")"
sm['A9'], sm['B9'] = 'Pass rate (of executed)', f"=IF(B8=0,0,COUNTIF('Test Cases'!N2:N{last},\"Pass\")/B8)"
sm['A10'], sm['B10'] = 'Known defects logged', f"=COUNTA('Known Defects'!A2:A{len(defects)+1})"
for r in range(4, 11):
    sm[f'A{r}'].font = bold
    sm[f'B{r}'].font = body_font
    sm[f'A{r}'].border = sm[f'B{r}'].border = border
sm['B7'].number_format = '0.0%'
sm['B9'].number_format = '0.0%'

def block(start_row, title, key_col, keys):
    sm.cell(row=start_row, column=1, value=title).font = Font(name=ARIAL, bold=True, size=11, color='1F3864')
    heads = ['', 'Total', 'Automated', 'Not Run', 'Pass', 'Fail', 'Blocked', 'N/A']
    for i, h in enumerate(heads, 1):
        c = sm.cell(row=start_row + 1, column=i, value=h if i > 1 else title.split(' by ')[-1].title())
        c.font, c.fill, c.border, c.alignment = hdr_font, hdr_fill, border, center
    r = start_row + 2
    rng = f"'Test Cases'!{key_col}2:{key_col}{last}"
    for k in keys:
        sm.cell(row=r, column=1, value=k)
        sm.cell(row=r, column=2, value=f'=COUNTIF({rng},A{r})')
        sm.cell(row=r, column=3, value=f"=COUNTIFS({rng},A{r},'Test Cases'!L2:L{last},\"Automated\")")
        for j, st in enumerate(['Not Run', 'Pass', 'Fail', 'Blocked', 'N/A'], 4):
            sm.cell(row=r, column=j, value=f"=COUNTIFS({rng},A{r},'Test Cases'!N2:N{last},\"{st}\")")
        for j in range(1, 9):
            c = sm.cell(row=r, column=j)
            c.font, c.border = body_font, border
            if j > 1: c.alignment = center
        r += 1
    sm.cell(row=r, column=1, value='Total').font = bold
    for j in range(2, 9):
        col = get_column_letter(j)
        c = sm.cell(row=r, column=j, value=f'=SUM({col}{start_row+2}:{col}{r-1})')
        c.font, c.border, c.alignment = bold, border, center
    sm.cell(row=r, column=1).border = border
    return r + 2

modules = list(dict.fromkeys(c[0] for c in cases))
types = sorted(set(c[8] for c in cases))
nr = block(12, 'By Module', 'B', modules)
nr = block(nr, 'By Type', 'J', types)
nr = block(nr, 'By Priority', 'K', ['P1', 'P2', 'P3'])
sm.column_dimensions['A'].width = 34
for col in 'BCDEFGH':
    sm.column_dimensions[col].width = 11

leg = nr
sm.cell(row=leg, column=1, value='How to use').font = Font(name=ARIAL, bold=True, size=11, color='1F3864')
tips = [
    'Yellow columns on "Test Cases" (Status, Actual Result, Defect ID / Remarks, Tested By, Test Date) are for you to fill in.',
    'Status is a dropdown: Not Run / Pass / Fail / Blocked / N/A. Colours update automatically.',
    '"Automated In" names the Playwright spec that runs the case: npx playwright test <spec> (see README).',
    'Red text in Expected Result = a KNOWN DEFECT already found on staging (see "Known Defects").',
    'Priority: P1 = must pass before release, P2 = should pass, P3 = nice to have.',
    'Example filled row: Status = Fail, Actual Result = "Footer link Forms opens 404", Defect ID = D-06, Tested By = your name, Test Date = 27-09-2026.',
]
for i, t in enumerate(tips, 1):
    c = sm.cell(row=leg + i, column=1, value=f'• {t}')
    c.font = body_font

# ---- Known Defects sheet
kd = wb.create_sheet('Known Defects')
kh = ['Defect ID', 'Category', 'Area', 'Description', 'Pages affected', 'Severity', 'URL', 'Found on', 'Status', 'Assigned To', 'Fixed In Build', 'Retest Result']
kw = [10, 16, 28, 60, 26, 10, 40, 12, 11, 16, 14, 14]
kd.append(kh)
for i, w in enumerate(kw, 1):
    kd.column_dimensions[get_column_letter(i)].width = w
    c = kd.cell(row=1, column=i)
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
for r, d in enumerate(defects, 2):
    did, cat, area, desc, pages, sev, url = d
    kd.append([did, cat, area, desc, pages, sev, BASE + url if url.startswith('/') else url, '27-09-2026', 'Open', '', '', ''])
    for ci in range(1, len(kh) + 1):
        c = kd.cell(row=r, column=ci)
        c.font, c.border = body_font, border
        c.alignment = center if ci in (1, 6, 8, 9) else wrap
        if ci >= 9: c.fill = input_fill
dv2 = DataValidation(type='list', formula1='"Open,In Progress,Fixed,Retest,Closed,Rejected"', allow_blank=True)
kd.add_data_validation(dv2)
dv2.add(f'I2:I{len(defects)+1}')
kd.freeze_panes = 'B2'
kd.auto_filter.ref = f'A1:L{len(defects)+1}'

# ---- Page Inventory sheet
pi = wb.create_sheet('Page Inventory')
ph = ['#', 'Page URL', 'Template', 'Interactive features found', 'Page-health test', 'Auto-click test', 'Manual review done?']
pw = [5, 52, 30, 50, 14, 14, 16]
pi.append(ph)
for i, w in enumerate(pw, 1):
    pi.column_dimensions[get_column_letter(i)].width = w
    c = pi.cell(row=1, column=i)
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
for r, (path, tpl, feats) in enumerate(inventory, 2):
    pi.append([r - 1, BASE + path, tpl, feats, 'Automated', 'Automated', ''])
    for ci in range(1, len(ph) + 1):
        c = pi.cell(row=r, column=ci)
        c.font, c.border = body_font, border
        c.alignment = center if ci in (1, 5, 6, 7) else wrap
        if ci == 7: c.fill = input_fill
        if ci == 4 and ('404' in feats):
            c.font = Font(name=ARIAL, size=10, bold=True, color='C00000')
dv3 = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True)
pi.add_data_validation(dv3)
dv3.add(f'G2:G{len(inventory)+1}')
n_inv = len(inventory) + 1
pi.cell(row=n_inv + 2, column=2, value='Menu/footer pages listed here; detail pages (≈60,000 press releases etc.) are sampled by the crawler (MAX_PER_SECTION).').font = Font(name=ARIAL, italic=True, size=9, color='595959')
pi.freeze_panes = 'C2'
pi.auto_filter.ref = f'A1:G{n_inv}'

for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 90

wb.save(OUT)
print(f'{len(cases)} test cases, {len(defects)} defects, {len(inventory)} pages -> {OUT}')
